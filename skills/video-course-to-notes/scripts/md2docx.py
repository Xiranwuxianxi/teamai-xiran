import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_font(run, size=None, bold=None, name_east='宋体', name_latin='Times New Roman'):
    run.font.name = name_latin
    r = run._element.rPr
    rFonts = r.find(qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), name_east)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold

def add_rich(p, text, size=10.5):
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok: continue
        if tok.startswith('**'):
            run = p.add_run(tok[2:-2]); set_font(run, size, True)
        elif tok.startswith('`'):
            run = p.add_run(tok[1:-1]); run.font.name='Courier New'; run.font.size=Pt(size)
        else:
            run = p.add_run(tok); set_font(run, size)

def convert(md_path, out_path, title_size=16):
    lines = open(md_path, encoding='utf-8').read().splitlines()
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(3.17)
        s.top_margin = s.bottom_margin = Cm(2.54)
    st = doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(10.5)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln or ln == '---':
            i += 1; continue
        if ln.startswith('|'):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl.append(lines[i].strip()); i += 1
            rows = []
            for t in tbl:
                if re.match(r'^\|[\s\-:|]+\|$', t): continue
                rows.append([c.strip() for c in t.strip('|').split('|')])
            if rows:
                tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
                tb.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        cp = tb.cell(ri, ci).paragraphs[0]
                        run = cp.add_run(re.sub(r'\*\*', '', cell))
                        set_font(run, 9.5, bold=(ri == 0))
            continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading('', level=min(lvl, 4))
            run = h.add_run(m.group(2))
            sizes = {1: title_size, 2: 14, 3: 12, 4: 11}
            set_font(run, sizes.get(lvl, 11), True, name_east='黑体', name_latin='Times New Roman')
            run.font.color.rgb = RGBColor(0, 0, 0)
            if lvl == 1: h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if re.match(r'^\s*[-*]\s+', ln):
            p = doc.add_paragraph(style='List Bullet'); add_rich(p, re.sub(r'^\s*[-*]\s+', '', ln)); i += 1; continue
        if re.match(r'^\s*\d+\.\s+', ln):
            p = doc.add_paragraph(style='List Number'); add_rich(p, re.sub(r'^\s*\d+\.\s+', '', ln)); i += 1; continue
        if ln.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                p = doc.add_paragraph(); run = p.add_run(lines[i]); run.font.name='Courier New'; run.font.size=Pt(9); i += 1
            i += 1; continue
        if ln.startswith('>'):
            p = doc.add_paragraph(); add_rich(p, ln.lstrip('> ')); p.paragraph_format.left_indent = Cm(0.75); i += 1; continue
        p = doc.add_paragraph()
        add_rich(p, ln)
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        i += 1
    doc.save(out_path)
    print('saved', out_path)

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
