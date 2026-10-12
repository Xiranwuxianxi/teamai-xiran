#!/usr/bin/env python3
"""Markdown -> .docx converter for CJK academic papers, no pandoc needed.
Usage: python3 md_to_docx_cjk.py input.md output.docx
Handles: # headings (黑体), **bold**, `code`, | pipe tables (real Word tables,
header row bold), bullet/numbered lists, ``` code blocks, > quotes.
Body: 宋体+Times New Roman 10.5pt, 1.5 line spacing, first-line indent, A4.
Requires: pip install python-docx
"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_font(run, size=None, bold=None, name_east='宋体', name_latin='Times New Roman'):
    run.font.name = name_latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name_east)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold

def add_rich(p, text, size=10.5):
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok: continue
        if tok.startswith('**'):
            set_font(p.add_run(tok[2:-2]), size, True)
        elif tok.startswith('`'):
            r = p.add_run(tok[1:-1]); r.font.name = 'Courier New'; r.font.size = Pt(size)
        else:
            set_font(p.add_run(tok), size)

def convert(md_path, out_path, title_size=16):
    lines = open(md_path, encoding='utf-8').read().splitlines()
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(3.17)
        s.top_margin = s.bottom_margin = Cm(2.54)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln or ln == '---':
            i += 1; continue
        if ln.startswith('|'):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in t.strip('|').split('|')]
                    for t in tbl if not re.match(r'^\|[\s\-:|]+\|$', t)]
            if rows:
                tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        set_font(tb.cell(ri, ci).paragraphs[0].add_run(re.sub(r'\*\*', '', cell)),
                                 9.5, bold=(ri == 0))
            continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading('', level=min(lvl, 4))
            run = h.add_run(m.group(2))
            set_font(run, {1: title_size, 2: 14, 3: 12, 4: 11}.get(lvl, 11), True,
                     name_east='黑体')
            run.font.color.rgb = RGBColor(0, 0, 0)
            if lvl == 1: h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if re.match(r'^\s*[-*]\s+', ln):
            add_rich(doc.add_paragraph(style='List Bullet'), re.sub(r'^\s*[-*]\s+', '', ln)); i += 1; continue
        if re.match(r'^\s*\d+\.\s+', ln):
            add_rich(doc.add_paragraph(style='List Number'), re.sub(r'^\s*\d+\.\s+', '', ln)); i += 1; continue
        if ln.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                r = doc.add_paragraph().add_run(lines[i]); r.font.name = 'Courier New'; r.font.size = Pt(9); i += 1
            i += 1; continue
        if ln.startswith('>'):
            p = doc.add_paragraph(); add_rich(p, ln.lstrip('> ')); p.paragraph_format.left_indent = Cm(0.75); i += 1; continue
        p = doc.add_paragraph(); add_rich(p, ln)
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        i += 1
    doc.save(out_path)
    print('saved', out_path)

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
